"""Read only the official label worksheet and verify every source video."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path


def sha256(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def write_json(path, data):
    path = Path(path)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    temporary.replace(path)


def read_samples(data_root):
    from openpyxl import load_workbook

    root = Path(data_root) / "attachment_1_raw_samples"
    wb = load_workbook(root / "labels_100.xlsx", read_only=True, data_only=True)
    try:
        rows = wb["label"].iter_rows(values_only=True)
        header = list(next(rows))
        needed = ["video_id", "clip_id", "text", "label", "annotation"]
        if not set(needed).issubset(header):
            raise ValueError(f"Unexpected label columns: {header}")
        samples, seen = [], set()
        for row in rows:
            if all(v is None for v in row):
                continue
            d = dict(zip(header, row))
            video_id = str(d["video_id"])
            clip = d["clip_id"]
            clip_id = str(int(clip)) if isinstance(clip, (int, float)) and float(clip).is_integer() else str(clip)
            key = (video_id, clip_id)
            if key in seen:
                raise ValueError(f"Duplicate label key: {key}")
            seen.add(key)
            y = float(d["label"])
            expected = "Negative" if y < 0 else "Positive" if y > 0 else "Neutral"
            if not -3 <= y <= 3 or d["annotation"] != expected:
                raise ValueError(f"Inconsistent label for {key}")
            path = root / "videos" / video_id / f"{clip_id}.mp4"
            if not path.is_file() or not str(d["text"] or "").strip():
                raise ValueError(f"Missing video or transcript: {key}")
            samples.append(dict(id=f"{video_id}$_${clip_id}", video_id=video_id, clip_id=clip_id,
                                text=str(d["text"]), label=y, annotation=expected, path=path))
        actual = {(p.parent.name, p.stem) for p in (root / "videos").glob("*/*.mp4")}
        if seen != actual or len(samples) != 100:
            raise ValueError(f"Expected 100 exact video/label pairs: labels={len(samples)}, videos={len(actual)}, mismatch={seen ^ actual}")
        return sorted(samples, key=lambda s: (s["video_id"], int(s["clip_id"])))
    finally:
        wb.close()
